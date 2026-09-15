"""
產生內建的 docxtpl 預設範本（vuln_weaver/templates/*.docx）。

範本就是一般 Word 檔，段落裡放 Jinja 標籤；使用者可以拿這兩個檔案當起點
自行調整版面。修改本腳本後重新執行即可重建範本：

    python scripts/build_default_template.py
"""
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent.parent / "vuln_weaver" / "templates"
PRIMARY = RGBColor(0x1F, 0x49, 0x7D)


def _margins(doc):
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1.0)
        section.left_margin = section.right_margin = Inches(1.0)


def _title(doc, text, size=24):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = PRIMARY
    return p


def _center(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def _heading(doc, text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    run.font.bold = True
    run.font.color.rgb = PRIMARY


def _para(doc, text):
    return doc.add_paragraph(text)


def _table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
            if r == 0:
                for run in table.rows[r].cells[c].paragraphs[0].runs:
                    run.font.bold = True
    return table


def _signature_table(doc):
    _table(doc, [
        ["執行人員", "審核人員", "核定主管"],
        ["姓名：{{ meta.tester }}", "姓名：{{ meta.reviewer }}", "姓名：{{ meta.approver }}"],
        ["簽章：\n\n\n", "簽章：\n\n\n", "簽章：\n\n\n"],
        ["日期：　　年　　月　　日", "日期：　　年　　月　　日", "日期：　　年　　月　　日"],
    ])


def build_report_template():
    doc = docx.Document()
    _margins(doc)
    doc.add_paragraph().paragraph_format.space_before = Pt(80)
    _title(doc, "資訊系統弱點掃描與安全健診報告書")
    _center(doc, "專案標的：{{ scan_name }}", 14)
    doc.add_paragraph()
    _table(doc, [
        ["項目", "內容"],
        ["專案／案號", "{{ meta.project_code }}"],
        ["受測單位", "{{ meta.org }}"],
        ["執行單位", "{{ meta.vendor }}"],
        ["掃描工具", "{{ scanner_label }}"],
        ["受測主機總數", "{{ host_count }} 台"],
        ["檢測日期", "{{ scan_date }}"],
        ["報告產出日期", "{{ generated_at }}"],
    ])
    doc.add_page_break()

    _heading(doc, "壹、 檢測作業與風險統計概況")
    _para(doc, "本次弱點掃描作業針對 {{ host_count }} 處目標主機進行自動化安全檢測，檢測日期為 {{ scan_date_iso }}，"
               "共彙整 {{ vuln_count }} 項獨立弱點。檢測結果分級統計如下：")
    _table(doc, [
        ["風險等級", "檢出數量", "建議修復時限"],
        ["極高 (Critical)", "{{ stats.Critical }}", "立即停用/隔離，並於 14 日內修復完成"],
        ["高 (High)", "{{ stats.High }}", "列入優先排程，並於 30 日內完成修補"],
        ["中 (Medium)", "{{ stats.Medium }}", "評估業務影響，於 60 日內定期維護修復"],
        ["低 (Low)", "{{ stats.Low }}", "於 90 日或次期系統升級時改善"],
        ["資訊 (Info)", "{{ stats.Info }}", "列為參考資訊，供管理人員監控"],
    ])
    doc.add_paragraph()
    _center(doc, "{{ chart }}")
    _center(doc, "圖 1.1：弱點風險等級分佈比例圖", 9)

    _heading(doc, "貳、 受檢主機資產與清冊")
    # docxtpl 的 {%tr %} 標籤要各自佔一整列：for 一列、資料一列、endfor 一列
    _table(doc, [
        ["序號", "主機 IP 位址", "電腦名稱 (FQDN)", "作業系統", "開放通訊埠"],
        ["{%tr for h in hosts %}", "", "", "", ""],
        ["{{ h.index }}", "{{ h.ip }}", "{{ h.hostname }}", "{{ h.os }}", "{{ h.open_ports }}"],
        ["{%tr endfor %}", "", "", "", ""],
    ])

    _heading(doc, "參、 弱點詳細檢出清冊與修補處置建議")
    _para(doc, "{%p if not vulns %}")
    _para(doc, "本檢測標的未檢出任何已知之嚴重、高、中或低等級弱點，安全防護狀態良好。")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p for v in vulns %}")
    p = doc.add_paragraph()
    run = p.add_run("項次 {{ v.index }}. [{{ v.severity_zh }}風險] {{ v.display_title }}")
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = PRIMARY
    _para(doc, "CVE 編號：{{ v.cves }}　　CVSS 評分：{{ v.cvss }}")
    _para(doc, "受影響主機／服務：{{ v.affected_hosts }}")
    _para(doc, "【弱點成因與風險說明】\n{{ v.description }}")
    _para(doc, "【建議修復與處置對策】\n{{ v.solution }}")
    small = _para(doc, "掃描器原始項目名稱：{{ v.title }}（Plugin ID: {{ v.id }}）")
    for run in small.runs:
        run.font.size = Pt(8.5)
    _para(doc, "{%p endfor %}")

    _heading(doc, "肆、 報告審核與簽章")
    _para(doc, "本報告經下列人員執行、審核與核定後生效；簽章欄請以親簽或電子簽章方式填具。")
    _signature_table(doc)

    out = OUT_DIR / "default_tw.docx"
    doc.save(str(out))
    return out


def build_diff_template():
    doc = docx.Document()
    _margins(doc)
    doc.add_paragraph().paragraph_format.space_before = Pt(80)
    _title(doc, "弱點改善複測比對與成效驗證報告")
    _center(doc, "初測標的：{{ baseline_scan_name }}", 13)
    _center(doc, "複測標的：{{ rescan_name }}", 13)
    doc.add_paragraph()
    _table(doc, [
        ["項目", "內容"],
        ["專案／案號", "{{ meta.project_code }}"],
        ["受測單位", "{{ meta.org }}"],
        ["執行單位", "{{ meta.vendor }}"],
        ["比對驗證日期", "{{ comparison_date }}"],
        ["總比對項目數", "{{ total_count }} 個風險項目"],
    ])
    doc.add_page_break()

    _heading(doc, "壹、 複測改善成效彙整")
    _table(doc, [
        ["複測狀態", "統計數量", "狀態說明"],
        ["已修復 (Fixed)", "{{ fixed_count }}", "初測時存在，經單位改善後於複測已不再檢出"],
        ["未修復 (Open)", "{{ open_count }}", "初測時存在且複測持續檢出，需持續列管追蹤"],
        ["新發現 (New)", "{{ new_count }}", "初測未檢出，複測階段新增之風險，需重新評估修補"],
    ])
    doc.add_paragraph()

    _heading(doc, "貳、 複測弱點逐項改善對照清單")
    _table(doc, [
        ["序號", "弱點名稱", "主機／服務", "風險等級", "複測狀態"],
        ["{%tr for i in items %}", "", "", "", ""],
        ["{{ i.index }}", "{{ i.title }}", "{{ i.affected_hosts }}", "{{ i.severity_zh }}", "{{ i.status_zh }}"],
        ["{%tr endfor %}", "", "", "", ""],
    ])
    doc.add_paragraph()

    _heading(doc, "參、 複測結果審核與簽章")
    _para(doc, "本報告經下列人員執行、審核與核定後生效；簽章欄請以親簽或電子簽章方式填具。")
    _signature_table(doc)

    out = OUT_DIR / "default_diff_tw.docx"
    doc.save(str(out))
    return out


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for path in (build_report_template(), build_diff_template()):
        print(f"已產生範本：{path}")
