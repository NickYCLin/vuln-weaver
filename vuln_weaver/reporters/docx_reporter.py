import os
import tempfile
from pathlib import Path
from typing import Union, Dict, Any, Optional
from datetime import datetime

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import matplotlib
matplotlib.use("Agg")  # Headless backend
import matplotlib.pyplot as plt

from vuln_weaver.reporters.base import BaseReporter
from vuln_weaver.models import ScanReport, DiffReport, Severity, DiffStatus


class DocxReporter(BaseReporter):
    """Generates professional, government/enterprise compliant Word (.docx) vulnerability audit reports."""

    def __init__(self):
        # Color palette
        self.COLOR_PRIMARY = RGBColor(0x1F, 0x49, 0x7D)   # Dark Blue
        self.COLOR_SECONDARY = RGBColor(0x59, 0x59, 0x59) # Gray
        self.COLOR_CRITICAL = RGBColor(0xD9, 0x53, 0x4F)  # Dark Red
        self.COLOR_HIGH = RGBColor(0xED, 0x6C, 0x02)      # Orange
        self.COLOR_MEDIUM = RGBColor(0xF0, 0xAD, 0x4E)    # Yellow/Amber
        self.COLOR_LOW = RGBColor(0x02, 0x88, 0xD1)       # Blue
        self.COLOR_INFO = RGBColor(0x75, 0x75, 0x75)      # Gray

    def generate(self, report: ScanReport, output_path: Union[str, Path], **kwargs) -> Path:
        """Generate a complete vulnerability assessment report."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()
        self._set_page_margins(doc)

        # 1. Cover Page
        self._build_cover_page(doc, report)
        doc.add_page_break()

        # 2. Section 1: Executive Summary
        self._build_executive_summary(doc, report)

        # 3. Section 2: Host Inventory
        self._build_host_inventory(doc, report)

        # 4. Section 3: Detailed Vulnerability Findings
        self._build_detailed_findings(doc, report)

        doc.save(str(out_file))
        return out_file

    def generate_diff(self, diff_report: DiffReport, output_path: Union[str, Path], **kwargs) -> Path:
        """Generate a re-scan diff audit report."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()
        self._set_page_margins(doc)

        # 1. Diff Cover Page
        self._build_diff_cover(doc, diff_report)
        doc.add_page_break()

        # 2. Diff Summary Table & Chart
        self._build_diff_summary(doc, diff_report)

        # 3. Detailed Diff Items
        self._build_diff_items(doc, diff_report)

        doc.save(str(out_file))
        return out_file

    # --- Internal Builder Methods ---

    def _set_page_margins(self, doc: docx.Document):
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

    def _build_cover_page(self, doc: docx.Document, report: ScanReport):
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(80)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run("資訊系統弱點掃描與安全健診報告書")
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = self.COLOR_PRIMARY

        # Subtitle
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(140)
        run_sub = p_sub.add_run(f"專案標的：{report.scan_name}")
        run_sub.font.size = Pt(16)
        run_sub.font.color.rgb = self.COLOR_SECONDARY

        # Metadata table
        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_items = [
            ("受測主機總數：", f"{len(report.hosts)} 台主機 (IP)"),
            ("掃描引擎：", f"{report.scanner_name.capitalize()} 掃描工具"),
            ("檢測產出日期：", report.scan_date.strftime("%Y 年 %m 月 %d 日")),
            ("產製單位：", "VulnWeaver 自動化合規檢核系統"),
        ]
        for idx, (label, val) in enumerate(meta_items):
            row = table.rows[idx]
            c1, c2 = row.cells[0], row.cells[1]
            c1.text = label
            c2.text = val
            c1.paragraphs[0].runs[0].font.bold = True
            c1.paragraphs[0].runs[0].font.size = Pt(11)
            c2.paragraphs[0].runs[0].font.size = Pt(11)

    def _build_executive_summary(self, doc: docx.Document, report: ScanReport):
        self._add_heading(doc, "壹、 檢測作業與風險統計概況", level=1)

        p = doc.add_paragraph()
        p.add_run(
            f"本次弱點掃描作業方針對本專案規劃之 {len(report.hosts)} 處目標主機進行自動化安全檢測，"
            f"檢測時間為 {report.scan_date.strftime('%Y-%m-%d')}。"
            "依據行政院國家資通安全防護規範與業界常見資安標準，檢測結果分級統計如下表所示："
        )

        # Risk Summary Table
        stats = report.summary_stats
        table = doc.add_table(rows=6, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["風險等級", "檢出數量", "建議修復時限 (公部門/企業標準)"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_bg(cell, "1F497D")
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        row_data = [
            ("極高 (Critical)", str(stats["Critical"]), "立即停用/隔離，並於 14 日內修復完成"),
            ("高 (High)", str(stats["High"]), "列入優先排程，並於 30 日內完成修補"),
            ("中 (Medium)", str(stats["Medium"]), "評估業務影響，於 60 日內定期維護修復"),
            ("低 (Low)", str(stats["Low"]), "於 90 日或次期系統升級時改善"),
            ("資訊 (Info)", str(stats["Info"]), "列為參考資訊，供管理人員監控"),
        ]

        for r_idx, (sev_label, count_str, sla_str) in enumerate(row_data, start=1):
            row = table.rows[r_idx]
            row.cells[0].text = sev_label
            row.cells[1].text = count_str
            row.cells[2].text = sla_str
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx % 2 == 0:
                self._set_cell_bg(row.cells[0], "F2F2F2")
                self._set_cell_bg(row.cells[1], "F2F2F2")
                self._set_cell_bg(row.cells[2], "F2F2F2")

        doc.add_paragraph()

        # Generate and insert Pie Chart
        chart_img = self._generate_severity_chart(stats)
        if chart_img:
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(chart_img, width=Inches(5.0))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_cap.add_run("圖 1.1：弱點風險等級分佈比例圖")
            r_cap.font.size = Pt(9.5)
            r_cap.font.color.rgb = self.COLOR_SECONDARY
            try:
                os.remove(chart_img)
            except OSError:
                pass

        doc.add_paragraph()

    def _build_host_inventory(self, doc: docx.Document, report: ScanReport):
        self._add_heading(doc, "貳、 受檢主機資產與清冊", level=1)

        table = doc.add_table(rows=len(report.hosts) + 1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序號", "主機 IP 位址", "電腦名稱 (FQDN)", "作業系統", "開放通訊埠"]

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_bg(cell, "1F497D")
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for idx, host in enumerate(report.hosts, start=1):
            row = table.rows[idx]
            row.cells[0].text = str(idx)
            row.cells[1].text = host.ip
            row.cells[2].text = host.hostname or "-"
            row.cells[3].text = host.os or "未明確辨識"
            row.cells[4].text = ", ".join(map(str, host.open_ports[:8])) if host.open_ports else "無開放"

            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx % 2 == 0:
                for c in row.cells:
                    self._set_cell_bg(c, "F2F2F2")

        doc.add_paragraph()

    def _build_detailed_findings(self, doc: docx.Document, report: ScanReport):
        self._add_heading(doc, "參、 弱點詳細檢出清冊與修補處置建議", level=1)

        if not report.vulnerabilities:
            p_none = doc.add_paragraph()
            p_none.add_run("本檢測標的未檢出任何已知之嚴重、高、中或低等級弱點，安全防護狀態良好。")
            return

        for idx, v in enumerate(report.vulnerabilities, start=1):
            display_title = v.title_zh or v.title
            table = doc.add_table(rows=6, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Header row: Title + Severity
            header_cell = table.rows[0].cells[0]
            header_cell.merge(table.rows[0].cells[1])
            header_cell.text = f"項次 {idx}. [{v.severity.zh_tw}風險] {display_title}"
            header_cell.paragraphs[0].runs[0].font.bold = True
            header_cell.paragraphs[0].runs[0].font.size = Pt(11)
            header_cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Severity color for header
            bg_color = {
                Severity.CRITICAL: "D9534F",
                Severity.HIGH: "ED6C02",
                Severity.MEDIUM: "F0AD4E",
                Severity.LOW: "0288D1",
                Severity.INFO: "757575",
            }.get(v.severity, "1F497D")
            self._set_cell_bg(header_cell, bg_color)

            # Row 1: CVE & CVSS
            cve_str = ", ".join(v.cve_list) if v.cve_list else "無特定 CVE"
            cvss_str = f"{v.cvss_score:.1f}" if v.cvss_score is not None else "未提供"
            table.rows[1].cells[0].text = f"CVE 編號：{cve_str}"
            table.rows[1].cells[1].text = f"CVSS 評分：{cvss_str}"

            # Row 2: Affected Hosts
            affected_hosts_cell = table.rows[2].cells[0]
            affected_hosts_cell.merge(table.rows[2].cells[1])
            affected_hosts_cell.text = f"受影響主機清單：\n" + ", ".join(v.affected_hosts)

            # Row 3: Description
            desc_cell = table.rows[3].cells[0]
            desc_cell.merge(table.rows[3].cells[1])
            desc_text = v.description_zh or v.description or "暫無描述"
            desc_cell.text = f"【弱點成因與風險說明】：\n{desc_text}"

            # Row 4: Solution
            sol_cell = table.rows[4].cells[0]
            sol_cell.merge(table.rows[4].cells[1])
            sol_text = v.solution_zh or v.solution or "請諮詢軟體原廠以取得最新補丁更新。"
            sol_cell.text = f"【建議修復與處置對策】：\n{sol_text}"

            # Row 5: English Title reference
            ref_cell = table.rows[5].cells[0]
            ref_cell.merge(table.rows[5].cells[1])
            ref_cell.text = f"掃描器原始項目名稱：{v.title} (Plugin ID: {v.id})"
            ref_cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            ref_cell.paragraphs[0].runs[0].font.color.rgb = self.COLOR_SECONDARY

            doc.add_paragraph()

    # --- Diff Report Builders ---

    def _build_diff_cover(self, doc: docx.Document, diff_report: DiffReport):
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(80)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run("弱點改善複測比對與成效驗證報告")
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = self.COLOR_PRIMARY

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(140)
        run_sub = p_sub.add_run(f"初測標的：{diff_report.baseline_scan_name}  ➔  複測標的：{diff_report.rescan_name}")
        run_sub.font.size = Pt(14)
        run_sub.font.color.rgb = self.COLOR_SECONDARY

        table = doc.add_table(rows=3, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        items = [
            ("比對驗證日期：", diff_report.comparison_date.strftime("%Y 年 %m 月 %d 日")),
            ("總比對項目數：", f"{len(diff_report.items)} 個風險項目"),
            ("驗證引擎：", "VulnWeaver 智慧差異比對器 (Diff Engine)"),
        ]
        for idx, (label, val) in enumerate(items):
            r = table.rows[idx]
            r.cells[0].text = label
            r.cells[1].text = val
            r.cells[0].paragraphs[0].runs[0].font.bold = True

    def _build_diff_summary(self, doc: docx.Document, diff_report: DiffReport):
        self._add_heading(doc, "壹、 複測改善成效彙整", level=1)

        table = doc.add_table(rows=4, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["複測狀態", "統計數量", "狀態說明"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_bg(cell, "1F497D")
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        rows = [
            ("已修復 (Fixed)", str(diff_report.fixed_count), "初測時存在，經單位改善後於複測已不再檢出 (驗收通過)"),
            ("未修復 (Open)", str(diff_report.open_count), "初測時存在且複測持續檢出，需請單位持續列管追蹤"),
            ("新發現 (New)", str(diff_report.new_count), "初測未檢出，複測階段新增之風險，需重新評估修補"),
        ]
        for r_idx, (st_name, count_str, desc_str) in enumerate(rows, start=1):
            r = table.rows[r_idx]
            r.cells[0].text = st_name
            r.cells[1].text = count_str
            r.cells[2].text = desc_str
            r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx % 2 == 0:
                for c in r.cells:
                    self._set_cell_bg(c, "F2F2F2")

        doc.add_paragraph()

    def _build_diff_items(self, doc: docx.Document, diff_report: DiffReport):
        self._add_heading(doc, "貳、 複測弱點逐項改善對照清單", level=1)

        table = doc.add_table(rows=len(diff_report.items) + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["序號", "弱點名稱", "風險等級", "複測狀態"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            self._set_cell_bg(cell, "1F497D")
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for idx, item in enumerate(diff_report.items, start=1):
            r = table.rows[idx]
            r.cells[0].text = str(idx)
            r.cells[1].text = f"{item.title}\n主機／服務：{', '.join(item.affected_hosts)}"
            r.cells[2].text = item.severity.zh_tw
            
            st_text = {
                DiffStatus.FIXED: "✔ 已修復",
                DiffStatus.OPEN: "✖ 未修復",
                DiffStatus.NEW: "⚠ 新增",
            }.get(item.status, str(item.status.value))
            r.cells[3].text = st_text

            # Align
            r.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Colors for status
            if item.status == DiffStatus.FIXED:
                self._set_cell_bg(r.cells[3], "D4EDDA")  # Light green
            elif item.status == DiffStatus.OPEN:
                self._set_cell_bg(r.cells[3], "F8D7DA")  # Light red
            elif item.status == DiffStatus.NEW:
                self._set_cell_bg(r.cells[3], "FFF3CD")  # Light yellow

        doc.add_paragraph()

    # --- Helper Utilities ---

    def _add_heading(self, doc: docx.Document, text: str, level: int = 1):
        p = doc.add_heading(level=level)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = self.COLOR_PRIMARY

    def _set_cell_bg(self, cell, hex_color: str):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def _generate_severity_chart(self, stats: Dict[str, int]) -> Optional[str]:
        """Generate a donut pie chart image and return its temp filepath."""
        # Configure CJK font support for Windows and Linux
        plt.rcParams["font.sans-serif"] = [
            "Microsoft JhengHei",
            "SimHei",
            "PingFang TC",
            "Noto Sans CJK TC",
            "DejaVu Sans",
        ]
        plt.rcParams["axes.unicode_minus"] = False

        labels = []
        sizes = []
        colors = []
        color_map = {
            "Critical": "#D9534F",
            "High": "#ED6C02",
            "Medium": "#F0AD4E",
            "Low": "#0288D1",
            "Info": "#757575",
        }
        zh_labels = {
            "Critical": "極高 (Critical)",
            "High": "高 (High)",
            "Medium": "中 (Medium)",
            "Low": "低 (Low)",
            "Info": "資訊 (Info)",
        }

        for sev in ["Critical", "High", "Medium", "Low", "Info"]:
            val = stats.get(sev, 0)
            if val > 0:
                labels.append(f"{zh_labels[sev]}: {val}")
                sizes.append(val)
                colors.append(color_map[sev])

        if not sizes:
            return None

        fig, ax = plt.subplots(figsize=(6, 3.5), subplot_kw=dict(aspect="equal"))
        wedges, texts, autotexts = ax.pie(
            sizes,
            autopct="%1.1f%%",
            pctdistance=0.75,
            colors=colors,
            startangle=140,
            textprops=dict(color="black", fontsize=9),
        )

        # Draw inner circle for donut look
        centre_circle = plt.Circle((0, 0), 0.50, fc="white")
        fig.gca().add_artist(centre_circle)

        # Add legend
        ax.legend(
            wedges,
            labels,
            title="弱點等級分佈",
            loc="center left",
            bbox_to_anchor=(1, 0, 0.5, 1),
            prop={"family": "sans-serif", "size": 9},
        )

        plt.tight_layout()

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp_path = tmp.name
            plt.savefig(tmp_path, dpi=180, bbox_inches="tight")
            plt.close(fig)
            return tmp_path
