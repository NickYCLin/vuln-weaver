"""
Excel (.xlsx) 匯出：弱點清冊、逐主機明細與複測列管表。

Word 是交付用，Excel 則是拿來排修補進度、篩選與追蹤的；列管表刻意留了
負責單位、預計完成日、備註等空欄給承辦人填。
"""
from pathlib import Path
from typing import Any, List, Optional, Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vuln_weaver.models import DiffReport, DiffStatus, ReportMeta, ScanReport, Severity
from vuln_weaver.reporters.base import BaseReporter

HEADER_FILL = PatternFill("solid", fgColor="1F497D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEVERITY_FILL = {
    Severity.CRITICAL: PatternFill("solid", fgColor="F4CCCC"),
    Severity.HIGH: PatternFill("solid", fgColor="FCE5CD"),
    Severity.MEDIUM: PatternFill("solid", fgColor="FFF2CC"),
    Severity.LOW: PatternFill("solid", fgColor="D9EAF7"),
    Severity.INFO: PatternFill("solid", fgColor="EFEFEF"),
}
STATUS_FILL = {
    DiffStatus.FIXED: PatternFill("solid", fgColor="D4EDDA"),
    DiffStatus.OPEN: PatternFill("solid", fgColor="F8D7DA"),
    DiffStatus.NEW: PatternFill("solid", fgColor="FFF3CD"),
}
STATUS_ZH = {DiffStatus.FIXED: "已修復", DiffStatus.OPEN: "未修復", DiffStatus.NEW: "新發現"}
WRAP = Alignment(wrap_text=True, vertical="top")


class XlsxReporter(BaseReporter):
    """Write ScanReport / DiffReport to an .xlsx workbook."""

    def generate(
        self,
        report: ScanReport,
        output_path: Union[str, Path],
        meta: Optional[ReportMeta] = None,
        **kwargs,
    ) -> Path:
        meta = meta or ReportMeta()
        wb = Workbook()

        self._summary_sheet(wb.active, report, meta)
        self._host_sheet(wb.create_sheet("主機清冊"), report)
        self._vuln_sheet(wb.create_sheet("弱點清冊"), report)
        self._detail_sheet(wb.create_sheet("逐主機明細"), report)

        return self._save(wb, output_path)

    def generate_diff(
        self,
        diff_report: DiffReport,
        output_path: Union[str, Path],
        meta: Optional[ReportMeta] = None,
        **kwargs,
    ) -> Path:
        meta = meta or ReportMeta()
        wb = Workbook()

        ws = wb.active
        ws.title = "複測摘要"
        rows = [
            ("初掃標的", diff_report.baseline_scan_name),
            ("複掃標的", diff_report.rescan_name),
            ("比對日期", diff_report.comparison_date.strftime("%Y-%m-%d")),
            ("已修復 (Fixed)", diff_report.fixed_count),
            ("未修復 (Open)", diff_report.open_count),
            ("新發現 (New)", diff_report.new_count),
            ("比對項目總數", len(diff_report.items)),
        ]
        self._key_value_sheet(ws, rows + self._meta_rows(meta))

        self._tracking_sheet(wb.create_sheet("複測列管表"), diff_report)
        return self._save(wb, output_path)

    # --- sheets -------------------------------------------------------

    def _summary_sheet(self, ws: Worksheet, report: ScanReport, meta: ReportMeta) -> None:
        ws.title = "摘要"
        stats = report.summary_stats
        rows = [
            ("專案標的", report.scan_name),
            ("掃描工具", report.scanner_label),
            ("掃描器版本", report.scanner_version or ""),
            ("檢測日期", report.scan_date.strftime("%Y-%m-%d")),
            ("受測主機數", len(report.hosts)),
            ("獨立弱點數", len(report.vulnerabilities)),
            ("極高 (Critical)", stats["Critical"]),
            ("高 (High)", stats["High"]),
            ("中 (Medium)", stats["Medium"]),
            ("低 (Low)", stats["Low"]),
            ("資訊 (Info)", stats["Info"]),
        ]
        self._key_value_sheet(ws, rows + self._meta_rows(meta))

    def _host_sheet(self, ws: Worksheet, report: ScanReport) -> None:
        headers = ["序號", "主機 IP／站台", "主機名稱 (FQDN)", "作業系統", "開放通訊埠", "極高", "高", "中", "低", "資訊"]
        self._write_table(ws, headers, [
            [
                idx, host.ip, host.hostname or "", host.os or "",
                ", ".join(map(str, host.open_ports)),
                host.vuln_count.get("Critical", 0), host.vuln_count.get("High", 0),
                host.vuln_count.get("Medium", 0), host.vuln_count.get("Low", 0), host.vuln_count.get("Info", 0),
            ]
            for idx, host in enumerate(report.hosts, start=1)
        ], widths=[6, 24, 28, 24, 30, 6, 6, 6, 6, 6])

    def _vuln_sheet(self, ws: Worksheet, report: ScanReport) -> None:
        headers = [
            "序號", "風險等級", "弱點名稱 (繁中)", "掃描器原始名稱", "弱點 ID", "CVSS", "CVE", "CWE",
            "受影響數", "受影響主機／服務", "弱點說明", "修補建議", "參考資料",
        ]
        rows = []
        for idx, v in enumerate(report.vulnerabilities, start=1):
            rows.append([
                idx, v.severity.zh_tw, v.title_zh or v.title, v.title, v.id,
                v.cvss_score if v.cvss_score is not None else "",
                ", ".join(v.cve_list), ", ".join(v.cwe_list),
                len(v.affected_hosts), "\n".join(v.affected_hosts),
                v.description_zh or v.description, v.solution_zh or v.solution,
                "\n".join(v.references),
            ])
        self._write_table(ws, headers, rows, widths=[6, 10, 40, 36, 12, 8, 24, 14, 9, 34, 60, 60, 40])
        for row_idx, v in enumerate(report.vulnerabilities, start=2):
            ws.cell(row=row_idx, column=2).fill = SEVERITY_FILL[v.severity]

    def _detail_sheet(self, ws: Worksheet, report: ScanReport) -> None:
        headers = ["序號", "主機／服務", "風險等級", "弱點名稱 (繁中)", "弱點 ID", "CVE", "修補建議", "修補狀態", "負責單位", "預計完成日", "備註"]
        rows = []
        idx = 0
        for v in report.vulnerabilities:
            for target in v.affected_hosts:
                idx += 1
                rows.append([idx, target, v.severity.zh_tw, v.title_zh or v.title, v.id,
                             ", ".join(v.cve_list), v.solution_zh or v.solution, "", "", "", ""])
        self._write_table(ws, headers, rows, widths=[6, 30, 10, 40, 12, 24, 60, 10, 14, 12, 24])
        row_idx = 2
        for v in report.vulnerabilities:
            for _ in v.affected_hosts:
                ws.cell(row=row_idx, column=3).fill = SEVERITY_FILL[v.severity]
                row_idx += 1

    def _tracking_sheet(self, ws: Worksheet, diff_report: DiffReport) -> None:
        headers = ["序號", "複測狀態", "風險等級", "弱點名稱", "弱點 ID", "主機／服務", "修補建議", "負責單位", "預計完成日", "備註"]
        rows = [
            [idx, STATUS_ZH[item.status], item.severity.zh_tw, item.title, item.vuln_id,
             "\n".join(item.affected_hosts), item.solution or "", "", "", ""]
            for idx, item in enumerate(diff_report.items, start=1)
        ]
        self._write_table(ws, headers, rows, widths=[6, 10, 10, 40, 14, 34, 60, 14, 12, 24])
        for row_idx, item in enumerate(diff_report.items, start=2):
            ws.cell(row=row_idx, column=2).fill = STATUS_FILL[item.status]
            ws.cell(row=row_idx, column=3).fill = SEVERITY_FILL[item.severity]

    # --- helpers ------------------------------------------------------

    @staticmethod
    def _meta_rows(meta: ReportMeta) -> List[tuple]:
        rows = [
            ("專案／案號", meta.project_code or ""),
            ("受測單位", meta.org or ""),
            ("執行單位", meta.vendor or ""),
            ("執行人員", meta.tester or ""),
            ("審核人員", meta.reviewer or ""),
            ("核定主管", meta.approver or ""),
        ]
        rows.extend((key, value) for key, value in meta.extra.items())
        return rows

    @staticmethod
    def _key_value_sheet(ws: Worksheet, rows: Sequence[tuple]) -> None:
        for key, value in rows:
            ws.append([key, value])
        for cell in ws["A"]:
            cell.font = Font(bold=True)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 60

    @staticmethod
    def _write_table(ws: Worksheet, headers: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[int]) -> None:
        ws.append(list(headers))
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(list(row))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = WRAP
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"

    @staticmethod
    def _save(wb: Workbook, output_path: Union[str, Path]) -> Path:
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out_file))
        return out_file
