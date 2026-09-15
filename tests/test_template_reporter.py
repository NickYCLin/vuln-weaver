from pathlib import Path

import docx
import pytest
from click.testing import CliRunner

from vuln_weaver.cli import main
from vuln_weaver.comparator.diff import VulnerabilityComparator
from vuln_weaver.models import ReportMeta
from vuln_weaver.parsers.nessus import NessusParser
from vuln_weaver.parsers.zap import ZapParser
from vuln_weaver.reporters.docx_reporter import DocxReporter
from vuln_weaver.reporters.template_reporter import (
    DEFAULT_DIFF_TEMPLATE,
    DEFAULT_REPORT_TEMPLATE,
    TemplateReporter,
    build_context,
)

NESSUS_FIXTURE = Path(__file__).parent / "fixtures" / "sample.nessus"
ZAP_FIXTURE = Path(__file__).parent / "fixtures" / "sample_zap.xml"


def _all_text(path: Path) -> str:
    document = docx.Document(str(path))
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def test_default_templates_exist():
    assert DEFAULT_REPORT_TEMPLATE.exists()
    assert DEFAULT_DIFF_TEMPLATE.exists()


def test_default_template_renders_full_report(tmp_path):
    report = NessusParser().parse(NESSUS_FIXTURE)
    meta = ReportMeta(org="某某市政府", vendor="安全顧問公司", project_code="114-INFOSEC-001",
                      tester="王小明", reviewer="李大華", approver="陳主管")

    out = TemplateReporter().generate(report, tmp_path / "tpl.docx", meta=meta)
    text = _all_text(out)

    assert "{{" not in text and "{%" not in text
    assert "2026 Q1 Internal Security Audit" in text
    assert "某某市政府" in text and "114-INFOSEC-001" in text
    assert "姓名：王小明" in text and "姓名：李大華" in text and "姓名：陳主管" in text
    assert "192.168.10.50" in text and "web-app01.local" in text
    assert "伺服器支援過期之 TLS 1.0 通訊協定" in text
    assert "Plugin ID: 104743" in text
    assert "Tenable Nessus" in text
    assert len(docx.Document(str(out)).inline_shapes) == 1  # 圓餅圖


def test_default_template_without_meta_leaves_blank_signature(tmp_path):
    report = ZapParser().parse(ZAP_FIXTURE)
    out = TemplateReporter().generate(report, tmp_path / "zap_tpl.docx")
    text = _all_text(out)
    assert "姓名：\n" in text or "姓名：" in text
    assert "OWASP ZAP" in text
    assert "portal.example.gov.tw" in text


def test_custom_template_and_extra_vars(tmp_path):
    template = tmp_path / "custom.docx"
    document = docx.Document()
    document.add_paragraph("案號 {{ meta.project_code }}，委託人 {{ meta.client_contact }}")
    document.add_paragraph("{%p for v in vulns %}")
    document.add_paragraph("- {{ v.severity_zh }} / {{ v.display_title }} / {{ v.affected_hosts }}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("共 {{ stats.total }} 項，主機 {{ host_count }} 台")
    document.save(str(template))

    report = NessusParser().parse(NESSUS_FIXTURE)
    meta = ReportMeta(project_code="CASE-42", extra={"client_contact": "張承辦"})
    out = TemplateReporter(template_path=template).generate(report, tmp_path / "custom_out.docx", meta=meta)
    text = _all_text(out)

    assert "案號 CASE-42，委託人 張承辦" in text
    assert "中 / 伺服器支援過期之 TLS 1.0 通訊協定 / 192.168.10.50:443/tcp, 192.168.10.51:443/tcp" in text
    assert "共 3 項，主機 2 台" in text


def test_template_values_are_escaped(tmp_path):
    report = NessusParser().parse(NESSUS_FIXTURE)
    meta = ReportMeta(org="A&B <單位>")
    out = TemplateReporter().generate(report, tmp_path / "escaped.docx", meta=meta)
    assert "A&B <單位>" in _all_text(out)


def test_missing_template_raises(tmp_path):
    report = NessusParser().parse(NESSUS_FIXTURE)
    with pytest.raises(FileNotFoundError, match="範本"):
        TemplateReporter(template_path=tmp_path / "nope.docx").generate(report, tmp_path / "x.docx")


def test_default_diff_template(tmp_path):
    baseline = NessusParser().parse(NESSUS_FIXTURE)
    rescan = NessusParser().parse(NESSUS_FIXTURE)
    diff = VulnerabilityComparator.compare(baseline, rescan)
    out = TemplateReporter().generate_diff(diff, tmp_path / "diff_tpl.docx", meta=ReportMeta(reviewer="李大華"))
    text = _all_text(out)
    assert "{{" not in text
    assert "未修復" in text and "姓名：李大華" in text
    assert "192.168.10.50:443/tcp" in text


def test_build_context_shape():
    ctx = build_context(NessusParser().parse(NESSUS_FIXTURE))
    assert ctx["stats"]["total"] == 3
    assert ctx["vulns"][0]["index"] == 1
    assert set(ctx["meta"]) >= {"org", "vendor", "tester", "reviewer", "approver", "signers"}


def test_docx_reporter_includes_signature_block(tmp_path):
    report = NessusParser().parse(NESSUS_FIXTURE)
    meta = ReportMeta(org="某某市政府", tester="王小明")
    out = DocxReporter().generate(report, tmp_path / "sig.docx", meta=meta)
    text = _all_text(out)
    assert "報告審核與簽章" in text
    assert "受測單位：" in text and "某某市政府" in text
    assert "姓名：王小明" in text
    assert "核定主管" in text

    diff = VulnerabilityComparator.compare(report, report)
    out_diff = DocxReporter().generate_diff(diff, tmp_path / "sig_diff.docx", meta=meta)
    assert "複測結果審核與簽章" in _all_text(out_diff)


def test_cli_parse_with_template_and_meta(tmp_path):
    out = tmp_path / "cli_tpl.docx"
    result = CliRunner().invoke(main, [
        "parse", str(NESSUS_FIXTURE), "-o", str(out),
        "--template", str(DEFAULT_REPORT_TEMPLATE),
        "--org", "某某市政府", "--tester", "王小明", "--var", "client_contact=張承辦",
    ])
    assert result.exit_code == 0, result.output
    text = _all_text(out)
    assert "某某市政府" in text and "姓名：王小明" in text


def test_cli_diff_with_template(tmp_path):
    out = tmp_path / "cli_diff_tpl.docx"
    result = CliRunner().invoke(main, [
        "diff", str(NESSUS_FIXTURE), str(NESSUS_FIXTURE), "-o", str(out),
        "--template", str(DEFAULT_DIFF_TEMPLATE), "--reviewer", "李大華",
    ])
    assert result.exit_code == 0, result.output
    assert "姓名：李大華" in _all_text(out)


def test_cli_rejects_bad_var_format(tmp_path):
    result = CliRunner().invoke(main, ["parse", str(NESSUS_FIXTURE), "-o", str(tmp_path / "x.docx"), "--var", "novalue"])
    assert result.exit_code != 0
    assert "key=value" in result.output
