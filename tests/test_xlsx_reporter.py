from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook

from vuln_weaver.cli import main
from vuln_weaver.comparator.diff import VulnerabilityComparator
from vuln_weaver.models import ReportMeta
from vuln_weaver.parsers.nessus import NessusParser
from vuln_weaver.parsers.zap import ZapParser
from vuln_weaver.reporters.xlsx_reporter import XlsxReporter

FIX = Path(__file__).parent / "fixtures"


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_xlsx_report_sheets_and_rows(tmp_path):
    report = NessusParser().parse(FIX / "sample.nessus")
    meta = ReportMeta(org="某某市政府", tester="王小明", extra={"承辦": "張承辦"})
    out = XlsxReporter().generate(report, tmp_path / "report.xlsx", meta=meta)

    wb = load_workbook(str(out))
    assert wb.sheetnames == ["摘要", "主機清冊", "弱點清冊", "逐主機明細"]

    summary = dict((k, v) for k, v in _rows(wb["摘要"]))
    assert summary["專案標的"] == "2026 Q1 Internal Security Audit"
    assert summary["掃描工具"] == "Tenable Nessus"
    assert summary["受測主機數"] == 2 and summary["中 (Medium)"] == 1
    assert summary["受測單位"] == "某某市政府" and summary["執行人員"] == "王小明" and summary["承辦"] == "張承辦"

    hosts = _rows(wb["主機清冊"])
    assert hosts[0][:5] == ["序號", "主機 IP／站台", "主機名稱 (FQDN)", "作業系統", "開放通訊埠"]
    assert hosts[1][1] == "192.168.10.50" and hosts[1][2] == "web-app01.local"
    assert len(hosts) == 3

    vulns = _rows(wb["弱點清冊"])
    assert len(vulns) == 4
    tls = next(r for r in vulns[1:] if r[4] == "104743")
    assert tls[1] == "中" and tls[2] == "伺服器支援過期之 TLS 1.0 通訊協定"
    assert tls[5] == 7.4 and "CVE-2011-3389" in tls[6]
    assert tls[8] == 2 and "192.168.10.50:443/tcp" in tls[9]
    assert wb["弱點清冊"].freeze_panes == "A2"
    assert wb["弱點清冊"].auto_filter.ref.startswith("A1:")

    detail = _rows(wb["逐主機明細"])
    assert detail[0][7:] == ["修補狀態", "負責單位", "預計完成日", "備註"]
    # 3 個弱點、共 4 個 host×vuln 組合
    assert len(detail) == 1 + sum(len(v.affected_hosts) for v in report.vulnerabilities)
    assert all(r[7] is None for r in detail[1:])


def test_xlsx_diff_tracking_sheet(tmp_path):
    baseline = ZapParser().parse(FIX / "sample_zap.xml")
    xml = (FIX / "sample_zap.xml").read_text(encoding="utf-8")
    start = xml.index("<alertitem>\n        <pluginid>40018</pluginid>")
    end = xml.index("</alertitem>", start) + len("</alertitem>")
    rescan_file = tmp_path / "rescan.xml"
    rescan_file.write_text(xml[:start] + xml[end:], encoding="utf-8")
    diff = VulnerabilityComparator.compare(baseline, ZapParser().parse(rescan_file))

    out = XlsxReporter().generate_diff(diff, tmp_path / "diff.xlsx", meta=ReportMeta(reviewer="李大華"))
    wb = load_workbook(str(out))
    assert wb.sheetnames == ["複測摘要", "複測列管表"]
    summary = dict((k, v) for k, v in _rows(wb["複測摘要"]))
    assert summary["已修復 (Fixed)"] == 1 and summary["未修復 (Open)"] == 3 and summary["審核人員"] == "李大華"

    tracking = _rows(wb["複測列管表"])
    assert tracking[0][:6] == ["序號", "複測狀態", "風險等級", "弱點名稱", "弱點 ID", "主機／服務"]
    fixed = [r for r in tracking[1:] if r[1] == "已修復"]
    assert len(fixed) == 1 and fixed[0][4] == "40018" and fixed[0][5] == "portal.example.gov.tw:443/tcp"


def test_cli_xlsx_outputs_and_default_names(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    runner = CliRunner()

    result = runner.invoke(main, ["parse", str(FIX / "sample.nessus"), "-f", "xlsx"])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "report.xlsx").exists()
    assert "Excel 弱點清冊" in result.output

    result = runner.invoke(main, ["diff", str(FIX / "sample.nessus"), str(FIX / "sample.nessus"), "-f", "xlsx"])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "diff_report.xlsx").exists()

    # 沒指定 -o 時 docx 也走預設檔名
    result = runner.invoke(main, ["parse", str(FIX / "sample.nessus")])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "report.docx").exists()

    # 副檔名和格式不符要提醒
    result = runner.invoke(main, ["parse", str(FIX / "sample.nessus"), "-f", "xlsx", "-o", "out.docx"])
    assert result.exit_code == 0, result.output
    assert "提醒" in result.output


def test_cli_rejects_template_with_xlsx(tmp_path):
    from vuln_weaver.reporters.template_reporter import DEFAULT_REPORT_TEMPLATE
    result = CliRunner().invoke(main, ["parse", str(FIX / "sample.nessus"), "-f", "xlsx", "-o", str(tmp_path / "x.xlsx"), "-t", str(DEFAULT_REPORT_TEMPLATE)])
    assert result.exit_code != 0
    assert "--template" in result.output
